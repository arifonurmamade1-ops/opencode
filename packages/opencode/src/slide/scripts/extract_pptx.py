import sys
import os
import warnings
warnings.filterwarnings("ignore")
warnings.filterwarnings("ignore", category=DeprecationWarning)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _citations
from pptx import Presentation
from pathlib import Path

def extract_pptx(path: str):
    prs = Presentation(path)
    slide_data_list = []
    for i, slide in enumerate(prs.slides):
        slide_data = {"slide": i + 1, "title": "", "content": [], "notes": ""}
        if slide.shapes.title:
            slide_data["title"] = slide.shapes.title.text
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_data["content"].append(shape.text)
        if slide.has_notes_slide:
            slide_data["notes"] = slide.notes_slide.notes_text_frame.text
        slide_data_list.append(slide_data)

    outline = []
    full_text_parts = []
    for s in slide_data_list:
        outline.append(f"Slide {s['slide']}: {s['title'] or (s['content'][0] if s['content'] else '(Sem título)')}")
        full_text_parts.append(f"--- Slide {s['slide']} ---")
        if s['title']:
            full_text_parts.append(f"Title: {s['title']}")
        full_text_parts.append("Content:")
        for c in s['content']:
            full_text_parts.append(f"- {c}")
        if s['notes']:
            full_text_parts.append(f"Notes: {s['notes']}")
        full_text_parts.append("")
    full_text = "\n".join(full_text_parts)
    return outline, full_text

if __name__ == "__main__":
    path = sys.argv[1]
    out_file = sys.argv[2] if len(sys.argv) > 2 else "extracted.md"
    output_lines = [f"Extracting: {path}"]
    output_lines.append("\n=== OUTLINE (TABLE OF CONTENTS) ===")
    try:
        outline, full_text = extract_pptx(path)
        output_lines.append("\n".join(outline))
    except Exception as e:
        full_text = ""
        output_lines.append(f"Could not get outline: {e}")

    output_lines.append("\n\n=== FULL TEXT ===")
    output_lines.append(full_text)

    output_lines.append("\n\n=== CITATIONS (JSON) ===")
    citations = _citations.extract_citations(full_text)
    output_lines.append(__import__("json").dumps(citations, ensure_ascii=False, indent=2))

    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(output_lines))
    print(f"Done. Output written to {out_file} | citations: {len(citations)}")
